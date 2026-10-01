import csv
import hashlib
import io
import zipfile
from itertools import islice
from datetime import datetime,date
from pathlib import Path
from django.core.files.base import ContentFile
from django.db import transaction
from django.shortcuts import get_object_or_404
from django.utils import timezone
from rest_framework import serializers
from rest_framework.decorators import api_view
from rest_framework.exceptions import ValidationError
from rest_framework.response import Response
from drf_spectacular.utils import extend_schema
from drf_spectacular.types import OpenApiTypes
from openpyxl import load_workbook
from crm.models import Branch,Source,Person,ContactMethod,ConsentRecord,StaffProfile
from crm.permissions import scope,profile,branch_scope
from crm.services import normalize_phone,duplicate_signals,next_reference,activity,audit,queue_candidates,notify
from crm.idempotency import idempotent
from .models import ImportBatch,ImportRow

FIELDS=['full_name','phone','email','preferred_country','preferred_course','temperature','lead_status','owner','source','legacy_id','consent']

def batch_data(batch):
    return {'id':str(batch.pk),'source_file':batch.source_file,'status':batch.status,'mapping':batch.mapping,'defaults':batch.defaults,'counts':batch.counts,'created_at':batch.created_at,'headers':list(batch.rows.first().raw) if batch.rows.exists() else []}

def read_file(file):
    ext=Path(file.name).suffix.lower()
    if file.size>10*1024*1024:raise ValidationError('Use a file of at most 10 MB.')
    raw=file.read()
    try:
        if ext=='.csv':
            content=raw.decode('utf-8-sig');reader=csv.DictReader(io.StringIO(content));headers=reader.fieldnames or [];rows=list(islice(reader,5001))
        elif ext=='.xlsx':
            with zipfile.ZipFile(io.BytesIO(raw)) as archive:
                if sum(item.file_size for item in archive.infolist())>64*1024*1024:raise ValidationError('The expanded workbook exceeds 64 MB.')
            book=load_workbook(io.BytesIO(raw),read_only=True,data_only=True,keep_links=False);sheet=book.active
            values=sheet.iter_rows(values_only=True);headers=[str(v or '').strip() for v in next(values)];rows=[]
            for values_row in values:
                if len(rows)>5000:break
                row={h:str(v if v is not None else '') if not isinstance(v,(date,datetime)) else v.isoformat() for h,v in zip(headers,values_row)}
                if any(row.values()):rows.append(row)
            book.close()
        else:raise ValidationError('Use an XLSX or UTF-8 CSV file. Save legacy XLS as XLSX first.')
    except (UnicodeError,ValueError,KeyError,StopIteration,OSError,zipfile.BadZipFile,csv.Error):raise ValidationError('Could not read the spreadsheet.')
    if len(rows)>5000 or not rows or not headers or len(headers)>100 or len(set(headers))!=len(headers) or any(not h for h in headers):raise ValidationError('Use 1–5,000 rows and unique, nonempty column headers (at most 100 columns).')
    if any(len(str(value))>10000 for row in rows for value in row.values()):raise ValidationError('A spreadsheet cell exceeds the text limit.')
    return raw,rows

@extend_schema(request=OpenApiTypes.OBJECT,responses=OpenApiTypes.OBJECT)
@api_view(['GET','POST'])
@transaction.atomic
def batches(request):
    scope(request.user,'import')
    if request.method=='GET':return Response([batch_data(b) for b in branch_scope(request.user,ImportBatch.objects.all(),'import').order_by('-created_at')[:100]])
    file=request.FILES.get('file')
    if not file:raise ValidationError('Upload an XLSX or CSV file.')
    branch=get_object_or_404(Branch,pk=request.data.get('branch_id',profile(request.user).branch_id))
    if scope(request.user,'import')!='full' and branch.pk!=profile(request.user).branch_id:raise ValidationError('Use your own branch.')
    raw,rows=read_file(file);fingerprint=hashlib.sha256(raw).hexdigest()
    batch,created=ImportBatch.objects.get_or_create(file_hash=fingerprint,branch=branch,defaults={'source_file':Path(file.name).name[:255],'created_by':request.user})
    if created:
        batch.file.save(f'{batch.pk}-{Path(file.name).name}',ContentFile(raw),save=True)
        ImportRow.objects.bulk_create([ImportRow(batch=batch,row_no=i+2,raw=row) for i,row in enumerate(rows)])
        audit(request,'IMPORT_STAGED',batch,new={'rows':len(rows),'source_file':batch.source_file})
    return Response({**batch_data(batch),'fields':FIELDS,'reused':not created},status=201 if created else 200)

@extend_schema(request=OpenApiTypes.OBJECT,responses=OpenApiTypes.OBJECT)
@api_view(['GET','POST'])
@transaction.atomic
def batch_action(request,pk,operation='preview'):
    batch=get_object_or_404(branch_scope(request.user,ImportBatch.objects.select_for_update(of=('self',)),'import'),pk=pk)
    if request.method=='GET':
        try:offset=max(0,int(request.query_params.get('offset',0)))
        except ValueError:raise ValidationError('Invalid offset.')
        return Response({**batch_data(batch),'total':batch.rows.count(),'offset':offset,'rows':[{'id':r.pk,'row_no':r.row_no,'raw':r.raw,'normalized':r.normalized,'result':r.result,'errors':r.errors,'candidates':r.candidates,'person_id':str(r.person_id) if r.person_id else None} for r in batch.rows.order_by('row_no')[offset:offset+100]],'fields':FIELDS})
    if operation=='validate':
        if batch.status not in ['STAGING','VALIDATED']:raise ValidationError('This batch cannot be revalidated.')
        mapping=request.data.get('mapping',{})
        defaults=request.data.get('defaults',{})
        if not isinstance(mapping,dict) or set(mapping)-set(FIELDS):raise ValidationError('Invalid column mapping.')
        headers=set(batch.rows.first().raw)
        if any(v not in headers for v in mapping.values() if v):raise ValidationError('Mapped columns must exist in the file.')
        if not mapping.get('full_name') or not mapping.get('phone'):raise ValidationError('Map the name and phone columns.')
        if not isinstance(defaults,dict):raise ValidationError('Defaults must be an object.')
        owner=get_object_or_404(StaffProfile,user_id=defaults.get('owner_id'),branch=batch.branch,role='COUNSELOR',user__is_active=True)
        source=get_object_or_404(Source,pk=defaults.get('source_id'),active=True)
        batch.mapping=mapping;batch.defaults={'owner_id':owner.user_id,'source_id':source.pk,'consent':serializers.BooleanField().run_validation(defaults.get('consent',False))}
        counts={'valid':0,'invalid':0,'duplicate':0,'needs_review':0};seen={}
        for row in batch.rows.order_by('row_no'):
            values={key:str(row.raw.get(column) or '').strip() for key,column in mapping.items() if column};errors=[];candidates=[]
            try:
                name=serializers.CharField(max_length=160).run_validation(values.get('full_name'))
                phone=normalize_phone(values.get('phone',''))
                email=serializers.EmailField(allow_blank=True).run_validation(values.get('email','')).lower()
                status=values.get('lead_status') or 'NEW';status=status.upper().replace(' ','_')
                temperature=(values.get('temperature') or 'WARM').upper()
                if status not in ['NEW','CONTACTED','COUNSELING','INTERESTED','DOCUMENT_COLLECTION','APPLICATION_READY']:raise ValidationError('Map unsupported LOST/ON_HOLD/CONVERTED states through reviewed CRM actions after import.')
                if temperature not in ['HOT','WARM','COLD']:raise ValidationError('Invalid temperature.')
                consent=batch.defaults['consent'] or values.get('consent','').lower() in ['yes','true','1']
                if not consent:raise ValidationError('Confirm documented consent before import.')
                owner_id=owner.user_id
                if values.get('owner'):
                    assigned=StaffProfile.objects.filter(branch=batch.branch,role='COUNSELOR',user__is_active=True,user__username=values['owner']).first()
                    if not assigned:raise ValidationError('Counselor username is not recognized in this branch.')
                    owner_id=assigned.user_id
                source_id=source.pk
                if values.get('source'):
                    source_item=Source.objects.filter(name__iexact=values['source'],active=True).first()
                    if not source_item:raise ValidationError('Source name is not recognized.')
                    source_id=source_item.pk
                matches=duplicate_signals(phone,email,name,country=values.get('preferred_country',''))
                candidates=[{'id':str(m['person'].pk),'confidence':m['confidence']} for m in matches if m['confidence']!='LOW']
                if phone in seen:candidates.append({'row_no':seen[phone],'confidence':'HIGH'})
                seen[phone]=row.row_no
                row.normalized={'full_name':name,'phone':phone,'email':email,'owner_id':owner_id,'source_id':source_id,'preferred_country':values.get('preferred_country','')[:80],'preferred_course':values.get('preferred_course','')[:160],'temperature':temperature,'lead_status':status,'legacy_id':values.get('legacy_id','')[:160]}
                row.result='DUPLICATE' if any(c['confidence'] in ['HIGH','EXACT'] for c in candidates) else 'NEEDS_REVIEW' if candidates else 'VALID'
            except ValidationError as e:row.result='INVALID';errors=[str(e.detail)];row.normalized={}
            row.errors=errors;row.candidates=candidates;row.save()
            counts[{'VALID':'valid','INVALID':'invalid','DUPLICATE':'duplicate','NEEDS_REVIEW':'needs_review'}[row.result]]+=1
        batch.counts=counts;batch.status='VALIDATED';batch.save();audit(request,'IMPORT_VALIDATED',batch,new=counts)
    elif operation=='review':
        if batch.status!='VALIDATED':raise ValidationError('Review rows before approving the batch.')
        from crm.operations import required_reason
        reason=required_reason(request.data)
        row=get_object_or_404(batch.rows.select_for_update(of=('self',)),pk=request.data.get('row_id'),result__in=['DUPLICATE','NEEDS_REVIEW'])
        decision=request.data.get('decision')
        if decision not in ['CREATE','SKIP']:raise ValidationError('Choose CREATE or SKIP after reviewing duplicate signals.')
        row.result='VALID' if decision=='CREATE' else 'SKIPPED';row.override_reason=reason;row.save()
        batch.counts={key:batch.rows.filter(result=value).count() for key,value in [('valid','VALID'),('invalid','INVALID'),('duplicate','DUPLICATE'),('needs_review','NEEDS_REVIEW'),('skipped','SKIPPED')]};batch.save()
        audit(request,'IMPORT_ROW_REVIEWED',batch,new={'row':row.row_no,'decision':decision,'reason':reason})
    elif operation=='approve':
        if batch.status!='VALIDATED':raise ValidationError('Validate the batch before approving.')
        if not batch.rows.filter(result='VALID').exists() and not (batch.rows.filter(result='COMMITTED').exists() and not batch.rows.filter(result__in=['DUPLICATE','NEEDS_REVIEW']).exists()):raise ValidationError('There are no valid rows to approve. Correct mapping or review duplicate rows first.')
        batch.status='APPROVED';batch.approved_by=request.user;batch.save();audit(request,'IMPORT_APPROVED',batch,new={'counts':batch.counts})
    elif operation=='rollback':
        if batch.status!='COMMITTED':raise ValidationError('Only committed batches can be rolled back.')
        from crm.operations import required_reason
        reason=required_reason(request.data)
        rows=list(batch.rows.filter(person__isnull=False).select_related('person'))
        changed=[r.row_no for r in rows if r.person.updated_at!=r.committed_person_updated_at or r.person.merged_into_id or r.person.archived_at]
        if changed:raise ValidationError({'rows':changed,'message':'Edited records require manual review; no records were archived.'})
        for row in rows:
            row.person.archived_at=timezone.now();row.person.save();audit(request,'IMPORT_PERSON_ROLLED_BACK',row.person,new={'batch_id':str(batch.pk),'reason':reason})
            row.result='ROLLED_BACK';row.save()
        batch.status='ROLLED_BACK';batch.save();audit(request,'IMPORT_ROLLED_BACK',batch,new={'reason':reason,'count':len(rows)})
    else:raise ValidationError('Unknown batch operation.')
    return Response(batch_data(batch))

@extend_schema(request=OpenApiTypes.OBJECT,responses=OpenApiTypes.OBJECT)
@api_view(['POST'])
@transaction.atomic
@idempotent(required=True)
def commit(request,pk):
    batch=get_object_or_404(branch_scope(request.user,ImportBatch.objects.select_for_update(of=('self',)),'import'),pk=pk)
    if batch.status=='COMMITTED':return Response(batch_data(batch))
    if batch.status!='APPROVED':raise ValidationError('A validated batch requires manager approval before commit.')
    count=0
    for row in batch.rows.filter(result='VALID',person__isnull=True).select_for_update(of=('self',)):
        data=dict(row.normalized);phone=data.pop('phone');email=data.pop('email');data.pop('legacy_id',None)
        # Recheck against writes since validation. Never silently merge or duplicate.
        matches=duplicate_signals(phone,email,data['full_name'],country=data.get('preferred_country',''))
        reviewed_ids={str(c.get('id')) for c in row.candidates if c.get('id')}
        reviewed_ids.update(str(value) for value in batch.rows.filter(row_no__in=[c['row_no'] for c in row.candidates if c.get('row_no')],person__isnull=False).values_list('person_id',flat=True))
        unreviewed=[m for m in matches if m['confidence']!='LOW' and (not row.override_reason or str(m['person'].pk) not in reviewed_ids)]
        if unreviewed:row.result='NEEDS_REVIEW';row.candidates=[{'id':str(m['person'].pk),'confidence':m['confidence']} for m in matches];row.save();continue
        person=Person.objects.create(**data,branch=batch.branch,ref=next_reference(),created_by=request.user,import_batch_id=batch.pk,import_row_number=row.row_no,source_file=batch.source_file)
        ContactMethod.objects.create(person=person,type='PHONE',raw_value=phone,normalized_value=phone,is_primary=True)
        if email:ContactMethod.objects.create(person=person,type='EMAIL',raw_value=email,normalized_value=email,is_primary=True)
        ConsentRecord.objects.create(person=person,recorded_by=request.user,method='IMPORT',wording_version='import-v1',metadata={'batch_id':str(batch.pk),'row':row.row_no})
        activity(request,person,'Imported from spreadsheet',f'{batch.source_file}, row {row.row_no}','IMPORT')
        from crm.calendar import start_sla
        start_sla(person,'FIRST_CONTACT')
        audit(request,'PERSON_IMPORTED',person,new={'batch_id':str(batch.pk),'row':row.row_no})
        person.refresh_from_db();row.person=person;row.result='COMMITTED';row.committed_person_updated_at=person.updated_at;row.save();count+=1
    # A conflict discovered after approval must remain reviewable, including partial commits.
    batch.status='VALIDATED' if batch.rows.filter(result__in=['DUPLICATE','NEEDS_REVIEW']).exists() else 'COMMITTED'
    batch.counts={key:batch.rows.filter(result=value).count() for key,value in [('valid','VALID'),('invalid','INVALID'),('duplicate','DUPLICATE'),('needs_review','NEEDS_REVIEW'),('skipped','SKIPPED'),('committed','COMMITTED')]}
    batch.save();audit(request,'IMPORT_COMMITTED',batch,new={'count':count,'status':batch.status})
    return Response(batch_data(batch))
